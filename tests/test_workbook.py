from openpyxl import Workbook as ExcelFile, load_workbook
from royalty_entry.workbook import Workbook, sheet_names


def test_reads_only_selected_sheet_columns_h_i_and_retains_identifiers(tmp_path):
    path = tmp_path/'source.xlsx'
    book = ExcelFile()
    book.active.title = 'Other'
    sheet = book.create_sheet('Payments')
    sheet['H1'], sheet['I1'] = 'Publisher', 'License'
    sheet['G2'], sheet['G3'], sheet['G4'] = 'PUB1', 'PUB2', 'PUB3'
    sheet['H2'], sheet['I2'] = 'Publisher A', '00123'
    sheet['H3'], sheet['I3'] = 'Publisher B', 456
    sheet['I3'].number_format = '000000'
    sheet['H4'], sheet['I4'] = 'Publisher C', '=I3'
    book.save(path)
    before = path.read_bytes()
    selected = Workbook(path, 'Payments', 2)
    assert [r.license for r in selected.records] == ['00123', '000456']
    assert selected.invalid[0][0] == 4
    assert sheet_names(path) == ['Other', 'Payments']
    copy = tmp_path/'copy.xlsx'
    selected.save_filtered(copy, set())
    assert path.read_bytes() == before
    assert load_workbook(copy)['Payments']['I2'].value == '00123'


def test_column_g_publisher_code_is_loaded_and_validated(tmp_path):
    path=tmp_path/'codes.xlsx'
    book=ExcelFile();sheet=book.active
    sheet['G2'],sheet['H2'],sheet['I2']='PUB832','Tre Ball','123'
    sheet['G3'],sheet['H3'],sheet['I3']='PB1395','Other','123'
    sheet['G4'],sheet['H4'],sheet['I4']='PUB83O','Bad code','123'
    book.save(path)
    selected=Workbook(path,'Sheet',2)
    assert [r.publisher_number for r in selected.records]==['PUB832','PB1395']
    assert selected.invalid[0][0]==4


def test_missing_code_is_review_even_when_name_matches(tmp_path):
    path=tmp_path/'missing-code.xlsx';book=ExcelFile();s=book.active
    s['H2'],s['I2']='Perfectly matching name','123'
    s['G3'],s['I3']='PUB832','123'
    book.save(path)
    result=Workbook(path,'Sheet',2)
    assert [r.row for r in result.records]==[3]
    assert result.records[0].publisher==''
    assert result.invalid[0][0]==2

