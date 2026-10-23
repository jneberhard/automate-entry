import csv
import shutil
from decimal import Decimal
from threading import Event
from openpyxl import Workbook
import pytest
from royalty_entry.engine import run, RunLock
from royalty_entry.journal import Journal
from royalty_entry.rules import Record


def test_live_run_preserves_partial_removals_and_source_after_failure(tmp_path, monkeypatch):
    source = tmp_path/'source.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Payments'
    sheet['H1'], sheet['I1'] = 'Publisher', 'License'
    for row, license_number in [(2, '123'), (3, '123'), (4, '456')]:
        sheet.cell(row, 8, 'Publisher')
        sheet.cell(row, 9, license_number)
    book.save(source)
    original = source.read_bytes()
    exports = []
    progress = []

    def export(self, destination, remove_rows, use_excel=False):
        exports.append((set(remove_rows), use_excel))
        shutil.copy2(self.path, destination)

    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_filtered', export)

    class InterruptedDriver:
        def open_record(self, record):
            if record.license == '456':
                raise RuntimeError('RDP disconnected')
            return True

        def details(self):
            return Decimal('-.02')

        def close_record(self):
            pass

    with pytest.raises(RuntimeError, match='RDP disconnected'):
        run(source, 'Payments', 2, 'Q3', True, tmp_path/'runs', InterruptedDriver, Event(), lambda _: None,
            progress=lambda *event: progress.append(event))
    assert exports == [(set(), True), ({2, 3}, False)]
    assert source.read_bytes() == original
    assert [(event[0], event[1]) for event in progress] == [(0, 3), (1, 3), (2, 3)]
    report = next((tmp_path/'runs').glob('*/report.csv'))
    with report.open(encoding='utf-8-sig', newline='') as file:
        statuses = [row[3] for row in list(csv.reader(file))[1:]]
    assert statuses == ['remove_balance', 'duplicate', 'STOPPED']
    journal = Journal(tmp_path/'runs'/'journal.sqlite3')
    assert journal.get(Record(2, '123', 'Publisher'), 'Q3')[0] == 'remove_balance'
    journal.close()


def test_run_lock_excludes_second_runner_and_releases(tmp_path):
    path = tmp_path/'run.lock'
    with RunLock(path):
        with pytest.raises(RuntimeError, match='Another royalty-entry run'):
            with RunLock(path):
                pytest.fail('Second run acquired the same lock')
    with RunLock(path):
        pass


def test_progress_includes_invalid_rows_duplicates_and_finalization(tmp_path, monkeypatch):
    source = tmp_path/'source.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Payments'
    sheet['H2'], sheet['I2'] = 'Publisher', '123'
    sheet['H3'], sheet['I3'] = 'Publisher', '123'
    sheet['H4'], sheet['I4'] = 'Missing license', None
    book.save(source)
    events = []

    class MissingDriver:
        def open_record(self, record):
            return False

    run(source, 'Payments', 2, 'Q3', False, tmp_path/'runs', MissingDriver, Event(), lambda _: None,
        progress=lambda *event: events.append(event))
    assert [(event[0], event[1]) for event in events] == [(0, 3), (1, 3), (2, 3), (3, 3), (3, 3)]
    assert 'Finalizing' in events[-1][2]


def test_saved_copy_keeps_shared_license_rows_and_report_explains_retention(tmp_path, monkeypatch):
    source = tmp_path/'source.xlsx'
    book = Workbook()
    sheet = book.active
    sheet['H1'], sheet['I1'] = 'Publisher', 'License'
    for row, license_number in [(2, '123'), (3, '456')]:
        sheet.cell(row, 8, 'Excluded')
        sheet.cell(row, 9, license_number)
    book.save(source)
    before = source.read_bytes()
    root = tmp_path/'runs'
    root.mkdir()
    journal = Journal(root/'journal.sqlite3')
    journal.put(Record(99, '123', 'Paid publisher'), 'Q3', 'paid', '-.10')
    journal.close()
    exports = []
    def export(self, destination, remove_rows, use_excel=False):
        exports.append(set(remove_rows))
        shutil.copy2(self.path, destination)
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_filtered', export)
    class ExcludedDriver:
        def open_record(self, record): return True
        def details(self): return Decimal('-1')
        def close_record(self): pass
    folder = run(source, 'Sheet', 2, 'Q3', True, root, ExcludedDriver, Event(), lambda _: None)
    assert exports == [set(), {3}]
    assert source.read_bytes() == before
    with (folder/'report.csv').open(encoding='utf-8-sig', newline='') as file:
        rows = list(csv.reader(file))
    assert rows[1][3] == 'keep_shared_license'
    assert rows[2][3] == 'remove_balance'


@pytest.fixture(autouse=True)
def isolate_results_tab_excel(monkeypatch):
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_results_tab',
                        lambda self, rows, name=None: name)


def test_live_run_writes_filtered_tab_and_preflights_before_driver(tmp_path, monkeypatch):
    source = tmp_path/'source.xlsx'
    book = Workbook()
    book.active['H2'], book.active['I2'] = 'Publisher', '123'
    book.save(source)
    calls = []
    def tab(self, rows, name=None):
        calls.append((set(rows), name))
        return name
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_results_tab', tab)
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_filtered',
                        lambda self, dest, rows, **kw: shutil.copy2(self.path, dest))
    class Driver:
        def open_record(self, record):
            assert calls == [(set(), None)]
            return True
        def details(self): return Decimal('-1')
        def close_record(self): pass
    logs = []
    run(source, 'Sheet', 2, 'Q3', True, tmp_path/'runs', Driver, Event(), logs.append)
    assert calls[1][0] == {2}
    assert calls[1][1].startswith('Royalty ')
    assert any(line.startswith('Results tab: Royalty ') for line in logs)


def test_locked_source_prevents_payments(tmp_path, monkeypatch):
    source = tmp_path/'source.xlsx'
    book = Workbook()
    book.active['H2'], book.active['I2'] = 'Publisher', '123'
    book.save(source)
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_filtered',
                        lambda self, dest, rows, **kw: shutil.copy2(self.path, dest))
    def locked(*args): raise RuntimeError('Save and close the workbook')
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_results_tab', locked)
    def forbidden(): pytest.fail('Driver started before workbook preflight')
    with pytest.raises(RuntimeError, match='Save and close'):
        run(source, 'Sheet', 2, 'Q3', True, tmp_path/'runs', forbidden, Event(), lambda _: None)
