from pathlib import Path
import os
import shutil
import sys
from .rules import Record, license_text, publisher_number


class Workbook:
    def __init__(self, path: Path, sheet: str, first_row: int):
        from openpyxl import load_workbook
        self.path = Path(path).resolve()
        if self.path.suffix.lower() not in {'.xlsx', '.xlsm'}:
            raise ValueError('Use an .xlsx or .xlsm workbook; save old .xls files as .xlsx first.')
        self.book = load_workbook(self.path, keep_vba=self.path.suffix.lower() == '.xlsm')
        if sheet not in self.book.sheetnames:
            raise ValueError('The selected worksheet does not exist')
        if first_row < 1:
            raise ValueError('First data row must be at least 1')
        self.sheet = self.book[sheet]
        self.sheet_name = sheet
        self.records = []
        self.invalid = []
        for row in range(first_row, self.sheet.max_row+1):
            publisher = self.sheet.cell(row, 8)
            code_cell = self.sheet.cell(row, 7)
            license_cell = self.sheet.cell(row, 9)
            if code_cell.value is None and publisher.value is None and license_cell.value is None:
                continue
            try:
                if license_cell.data_type == 'f':
                    raise ValueError('License cells cannot contain formulas')
                if code_cell.data_type == 'f':
                    raise ValueError('Publisher number cannot contain a formula')
                code = publisher_number(code_cell.value)
                license_number = license_text(license_cell.value, license_cell.number_format)
                self.records.append(Record(row, license_number, str(publisher.value or '').strip(), code))
            except ValueError as exc:
                self.invalid.append((row, str(exc)))

    def save_filtered(self, destination: Path, remove_rows: set[int], use_excel=False):
        if destination.resolve() == self.path:
            raise ValueError('The original workbook cannot be overwritten')
        temp = destination.with_name(destination.stem + '.tmp' + destination.suffix)
        # Excel itself adjusts formulas, tables, names and formatting after row deletion.
        # Only a copied file is ever opened for editing.
        shutil.copy2(self.path, temp)
        if remove_rows or use_excel:
            sys.coinit_flags = 0
            import pythoncom
            import win32com.client
            pythoncom.CoInitializeEx(pythoncom.COINIT_MULTITHREADED)
            excel = book = None
            try:
                excel = win32com.client.DispatchEx('Excel.Application')
                excel.Visible = False
                excel.DisplayAlerts = False
                excel.AskToUpdateLinks = False
                excel.AutomationSecurity = 3  # Never run workbook macros.
                book = excel.Workbooks.Open(str(temp.resolve()), UpdateLinks=0, ReadOnly=False)
                sheet = book.Worksheets(self.sheet_name)
                for row in sorted(remove_rows, reverse=True):
                    sheet.Rows(row).Delete()
                book.Save()
            finally:
                try:
                    if book is not None:
                        book.Close(SaveChanges=False)
                finally:
                    try:
                        if excel is not None:
                            excel.Quit()
                    finally:
                        pythoncom.CoUninitialize()
        os.replace(temp, destination)

    def save_results_tab(self, remove_rows, name=None):
        """Copy the source tab within Excel and save only after filtering succeeds.

        With no name, check write access before any payments; do not save.
        The workbook must be closed in other Excel instances.
        """
        sys.coinit_flags = 0
        import pythoncom
        import win32com.client
        pythoncom.CoInitializeEx(pythoncom.COINIT_MULTITHREADED)
        excel = book = None
        try:
            excel = win32com.client.DispatchEx('Excel.Application')
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AskToUpdateLinks = False
            excel.AutomationSecurity = 3
            excel.EnableEvents = False
            book = excel.Workbooks.Open(str(self.path), UpdateLinks=0, ReadOnly=False)
            if book.ReadOnly:
                raise RuntimeError('Save and close the workbook in Excel before running; it must be writable to add a results tab.')
            if book.ProtectStructure:
                raise RuntimeError('Workbook structure is protected; Excel cannot add a results tab.')
            if name is None:
                return
            existing = {str(book.Worksheets(i).Name).casefold() for i in range(1, book.Worksheets.Count+1)}
            base, candidate, suffix = name[:31], name[:31], 2
            while candidate.casefold() in existing:
                tail = f' ({suffix})'
                candidate = base[:31-len(tail)] + tail
                suffix += 1
            source = book.Worksheets(self.sheet_name)
            # Dynamic COM may invoke Copy as a property with no arguments,
            # creating a different workbook. Invoke its method explicitly.
            count = book.Worksheets.Count
            source._oleobj_.InvokeTypes(0x227, 0, pythoncom.DISPATCH_METHOD,
                (24, 0), ((12, 17), (12, 17)),
                pythoncom.Empty, book.Worksheets(count))
            if book.Worksheets.Count != count+1:
                raise RuntimeError('Excel did not create the results tab in the selected workbook; no changes saved.')
            result = book.Worksheets(book.Worksheets.Count)
            result.Name = candidate
            for row in sorted(remove_rows, reverse=True):
                result.Rows(row).Delete()
            book.Save()
            return candidate
        finally:
            try:
                if book is not None:
                    book.Close(SaveChanges=False)
            finally:
                try:
                    if excel is not None:
                        excel.Quit()
                finally:
                    pythoncom.CoUninitialize()


def sheet_names(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()
