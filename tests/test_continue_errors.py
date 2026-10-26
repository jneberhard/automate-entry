from decimal import Decimal
from threading import Event
from types import SimpleNamespace
import csv
import shutil
import pytest
from royalty_entry.engine import process_records, run
from royalty_entry.journal import Journal
from royalty_entry.rules import Record
from royalty_entry.desktop import RemoteDriver, ScreenError, Cancelled, PaymentNotSaved
from royalty_entry.workbook import ERROR_BLUE, highlight_errors

class Driver:
    def __init__(self, fail=None):
        self.fail=fail;self.opened=[];self.saved=[];self.recovered=0
    def open_record(self,record):self.current=record;self.opened.append(record.license);return True
    def details(self):
        if self.current.license=='111':raise ValueError('Unreadable balance')
        return Decimal('-.02') if self.current.license=='222' else Decimal('0')
    def transfers(self):return Decimal('-.25')
    def close_record(self,**kw):pass
    def prepare_payment(self,*a):pass
    def commit_payment(self,*a):
        if self.fail and self.current.license=='333':raise self.fail('Uncertain save')
        self.saved.append(self.current.license)
    def recover_after_error(self):self.recovered+=1

def test_errors_and_duplicates_skip_but_next_license_is_processed(tmp_path):
    j=Journal(tmp_path/'journal.sqlite3');d=Driver();rows=[]
    records=[Record(2,'111','A','PUB1'),Record(3,'111','A','PUB1'),Record(4,'333','B','PUB2')]
    removed=process_records(records,d,'Q3',j,True,lambda *a:rows.append(a),Event())
    assert removed==set() and d.saved==['333'] and d.opened==['111','333'] and d.recovered==1
    assert [row[1] for row in rows]==['error','duplicate','paid']
    assert 'Unreadable balance' in rows[1][3]
    j.close()

@pytest.mark.parametrize('failure,expected',[(RuntimeError,'pending'),(PaymentNotSaved,'not_saved')])
def test_failed_save_is_not_retried_and_other_license_can_continue(tmp_path,failure,expected):
    j=Journal(tmp_path/'journal.sqlite3');d=Driver(failure);rows=[]
    records=[Record(2,'333','A','PUB1'),Record(3,'333','A','PUB1'),Record(4,'444','B','PUB2')]
    process_records(records,d,'Q3',j,True,lambda *a:rows.append(a),Event())
    assert j.get(records[0],'Q3')[0]==expected
    assert d.saved==['444'] and d.opened==['333','444'] and d.recovered==1
    j.close()

def test_pending_record_is_reported_without_opening_then_next_is_processed(tmp_path):
    j=Journal(tmp_path/'journal.sqlite3');d=Driver();rows=[]
    records=[Record(2,'333','A','PUB1'),Record(3,'444','B','PUB2')]
    j.put(records[0],'Q3','pending','-.26')
    process_records(records,d,'Q3',j,True,lambda *a:rows.append(a),Event())
    assert d.opened==['444'] and d.saved==['444'] and not d.recovered
    assert rows[0][1]=='error' and j.get(records[0],'Q3')[0]=='pending'
    j.close()

def test_failed_recovery_stops_before_opening_the_next_record(tmp_path):
    j=Journal(tmp_path/'journal.sqlite3');d=Driver();rows=[]
    def recover():raise ScreenError('Remote Desktop lost focus')
    d.recover_after_error=recover
    with pytest.raises(RuntimeError,match='cannot safely return'):
        process_records([Record(2,'111','A','PUB1'),Record(3,'444','B','PUB2')],d,'Q3',j,True,lambda *a:rows.append(a),Event())
    assert d.opened==['111'] and rows[0][1]=='error'
    j.close()

def test_user_stop_is_not_treated_as_row_error(tmp_path):
    j=Journal(tmp_path/'journal.sqlite3');d=Driver();rows=[]
    def stopped():raise Cancelled('User stopped')
    d.details=stopped
    with pytest.raises(Cancelled):process_records([Record(2,'111','A','PUB1')],d,'Q3',j,True,lambda *a:rows.append(a),Event())
    assert not rows and not d.recovered
    j.close()

def test_dialog_recovery_only_sends_escape_and_verifies_manager(monkeypatch):
    from royalty_entry import desktop
    d=RemoteDriver.__new__(RemoteDriver);d.handle=1;d.record=object();keys=[]
    stages=iter(['chooser','entry','financials','details','manager','manager'])
    d.snapshot=lambda:next(stages);d.state=lambda image:image;d.check_stop=lambda:None;d.wait=lambda *a:None
    d.expect=lambda stage: assert_stage(next(stages),stage)
    d.hotkey=lambda *args:keys.append(args)
    monkeypatch.setattr(desktop,'foreground',lambda:(1,'Remote Desktop',None))
    d.recover_after_error()
    assert keys==[('escape',)]*4 and d.record is None

def assert_stage(actual,expected):
    assert actual==expected
    return actual

def test_unknown_dialog_never_receives_recovery_keys():
    d=RemoteDriver.__new__(RemoteDriver);d.check_stop=lambda:None;d.snapshot=lambda:None;d.state=lambda _: 'unknown'
    d.hotkey=lambda *a:pytest.fail('Unknown dialog received input')
    with pytest.raises(ScreenError):d.recover_after_error()

def test_highlight_marks_only_requested_rows():
    bands={}
    def band(a,b):
        assert a[0]==b[0] and a[1]==1 and b[1]==9
        result=SimpleNamespace(Interior=SimpleNamespace(Color=None),Font=SimpleNamespace(Color=None));bands[a[0]]=result;return result
    sheet=SimpleNamespace(Cells=lambda row,col:(row,col),Range=band)
    highlight_errors(sheet,{3,7},9)
    assert set(bands)=={3,7}
    assert all(item.Interior.Color==ERROR_BLUE and item.Font.Color==0 for item in bands.values())

@pytest.mark.parametrize('live',[True,False])
def test_run_collects_all_errors_highlights_output_and_preserves_source(tmp_path,monkeypatch,live):
    from openpyxl import Workbook
    source=tmp_path/'input.xlsx';book=Workbook();sheet=book.active
    for row,license_number in [(2,'111'),(3,'111'),(4,'222'),(5,'333')]:
        sheet.cell(row,7,'PUB1');sheet.cell(row,8,'Publisher');sheet.cell(row,9,license_number)
    sheet.cell(6,7,'PUB1');sheet.cell(6,8,'Missing license')
    book.save(source);before=source.read_bytes();exports=[];tabs=[]
    def export(self,destination,remove_rows,use_excel=False,error_rows=None):
        exports.append((set(remove_rows),set(error_rows or ())));shutil.copy2(self.path,destination)
    def tab(self,remove_rows,name=None,error_rows=None):tabs.append((set(remove_rows),set(error_rows or ())));return name
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_filtered',export)
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_results_tab',tab)
    d=Driver();logs=[]
    folder=run(source,'Sheet',2,'Q3',live,tmp_path/'runs',lambda:d,Event(),logs.append)
    with (folder/'errors.csv').open(encoding='utf-8-sig',newline='') as file:errors=list(csv.DictReader(file))
    assert {int(e['Original Excel row']) for e in errors}=={2,3,6}
    assert exports[-1]==({4} if live else set(),{2,3,6})
    assert tabs[-1]==({4},{2,3,6}) if live else not tabs
    assert source.read_bytes()==before
    assert len([line for line in logs if line.startswith('ERROR')])==3

def test_all_invalid_rows_are_retained_blue_without_starting_driver(tmp_path,monkeypatch):
    from openpyxl import Workbook
    source=tmp_path/'input.xlsx';book=Workbook();book.active['G2']='BAD';book.active['I2']='123';book.save(source)
    exports=[]
    def export(self,destination,remove_rows,use_excel=False,error_rows=None):
        exports.append(set(error_rows or ()));
        shutil.copy2(self.path,destination)
    monkeypatch.setattr('royalty_entry.workbook.Workbook.save_filtered',export)
    def forbidden():pytest.fail('Invalid rows must not start screen automation')
    folder=run(source,'Sheet',2,'Q3',False,tmp_path/'runs',forbidden,Event(),lambda _:None)
    assert exports[-1]=={2}
    with (folder/'errors.csv').open(encoding='utf-8-sig',newline='') as file:assert len(list(csv.DictReader(file)))==1

def test_gui_queues_all_row_errors_and_completion_summary(tmp_path,monkeypatch):
    from queue import Queue
    from royalty_entry.gui import App
    from royalty_entry import engine
    folder=tmp_path/'run';folder.mkdir()
    with (folder/'errors.csv').open('w',encoding='utf-8-sig',newline='') as file:
        writer=csv.writer(file)
        writer.writerow(['Original Excel row','License','Publisher code','Publisher','Error'])
        writer.writerow([2,'111','PUB1','Publisher','Unreadable balance'])
        writer.writerow([3,'222','PB1395','Publisher','Pending save'])
    monkeypatch.setattr(engine,'run',lambda *a,**k:folder)
    app=SimpleNamespace(messages=Queue(),stop=Event())
    App.worker(app,('file.xlsx','Sheet',2,'Q3',False),{},False)
    kind,(path,errors)=app.messages.get_nowait()
    assert kind=='run_errors' and len(errors)==2 and path.endswith('errors.csv')
    kind,message=app.messages.get_nowait()
    assert kind=='done' and '2 row(s)' in message and 'marked blue' in message
